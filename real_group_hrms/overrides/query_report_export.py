import datetime

from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from frappe.desk import query_report
from frappe.utils import cint, cstr
from frappe.utils.xlsxutils import ILLEGAL_CHARACTERS_RE, handle_html


DEFAULT_XLSX_FONT_NAME = "Calibri"
HEADER_LIKE_XLSX_FONT_SIZE = 13


def make_xlsx_style(*, bold=False, font_size=11, font_name=DEFAULT_XLSX_FONT_NAME):
    return {
        "bold": 1 if bold else 0,
        "font_size": font_size,
        "font_name": font_name,
    }


def get_header_like_xlsx_style():
    return make_xlsx_style(bold=True, font_size=HEADER_LIKE_XLSX_FONT_SIZE)


def ensure_xlsx_styles(row):
    styles = row.get("_xlsx_styles")
    if not styles:
        styles = {}
        row["_xlsx_styles"] = styles
    return styles


def set_xlsx_cell_style(row, fieldname, *, bold=False, font_size=11, font_name=DEFAULT_XLSX_FONT_NAME):
    styles = ensure_xlsx_styles(row)
    styles[fieldname] = make_xlsx_style(
        bold=bold,
        font_size=font_size,
        font_name=font_name,
    )
    return row


def set_xlsx_cell_style_like_header(row, fieldname):
    styles = ensure_xlsx_styles(row)
    styles[fieldname] = get_header_like_xlsx_style()
    return row


def set_xlsx_styles(row, fieldnames, *, bold=False, font_size=11, font_name=DEFAULT_XLSX_FONT_NAME):
    for fieldname in fieldnames:
        set_xlsx_cell_style(
            row,
            fieldname,
            bold=bold,
            font_size=font_size,
            font_name=font_name,
        )
    return row


def build_xlsx_data_with_cell_styles(
    data,
    visible_idx,
    include_indentation,
    include_filters=False,
    ignore_visible_idx=False,
):
    excel_types = (
        str,
        bool,
        type(None),
        int,
        float,
        datetime.datetime,
        datetime.date,
        datetime.time,
        datetime.timedelta,
    )

    if len(visible_idx) == len(data.result) or not visible_idx:
        ignore_visible_idx = True
    else:
        visible_idx = set(visible_idx)

    result = []
    column_widths = []

    if cint(include_filters):
        filter_data = []
        filters = data.filters
        for filter_name, filter_value in filters.items():
            if not filter_value:
                continue
            filter_value = (
                ", ".join([cstr(x) for x in filter_value])
                if isinstance(filter_value, list)
                else cstr(filter_value)
            )
            filter_data.append([cstr(filter_name), filter_value])
        filter_data.append([])
        result += filter_data

    column_data = []
    for column in data.columns:
        if column.get("hidden"):
            continue
        column_data.append(query_report._(column.get("label")))
        column_width = cint(column.get("width", 0))
        column_width /= 10
        column_widths.append(column_width)
    result.append(column_data)

    for row_idx, row in enumerate(data.result):
        if not (ignore_visible_idx or row_idx in visible_idx):
            continue

        row_data = []
        if isinstance(row, dict):
            xlsx_styles = row.get("_xlsx_styles") or {}
            for col_idx, column in enumerate(data.columns):
                if column.get("hidden"):
                    continue

                label = column.get("label")
                fieldname = column.get("fieldname")
                cell_value = row.get(fieldname, row.get(label, ""))
                if not isinstance(cell_value, excel_types):
                    cell_value = cstr(cell_value)

                if cint(include_indentation) and "indent" in row and col_idx == 0:
                    cell_value = ("    " * cint(row["indent"])) + cstr(cell_value)

                style = xlsx_styles.get(fieldname)
                if style:
                    cell_value = {"value": cell_value, "style": style}

                row_data.append(cell_value)
        elif row:
            row_data = row

        result.append(row_data)

    return result, column_widths


def make_xlsx_with_cell_styles(data, sheet_name, wb=None, column_widths=None):
    from frappe.utils.xlsxutils import INVALID_TITLE_REGEX, get_column_letter, openpyxl
    from io import BytesIO

    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook(write_only=True)

    sheet_name_sanitized = INVALID_TITLE_REGEX.sub(" ", sheet_name)
    ws = wb.create_sheet(sheet_name_sanitized, 0)

    for i, column_width in enumerate(column_widths):
        if column_width:
            ws.column_dimensions[get_column_letter(i + 1)].width = column_width

    row1 = ws.row_dimensions[1]
    row1.font = Font(name="Calibri", bold=True)

    for row in data:
        clean_row = []
        for item in row:
            cell_style = None
            if isinstance(item, dict) and "value" in item:
                cell_style = item.get("style") or {}
                item = item.get("value")

            if isinstance(item, str) and (sheet_name not in ["Data Import Template", "Data Export"]):
                value = handle_html(item)
            else:
                value = item

            if isinstance(item, str) and next(ILLEGAL_CHARACTERS_RE.finditer(value), None):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)

            if cell_style:
                cell = WriteOnlyCell(ws, value=value)
                cell.font = Font(
                    name=cell_style.get("font_name", DEFAULT_XLSX_FONT_NAME),
                    bold=bool(cell_style.get("bold")),
                    size=cell_style.get("font_size", 11),
                )
                clean_row.append(cell)
            else:
                clean_row.append(value)

        ws.append(clean_row)

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file
